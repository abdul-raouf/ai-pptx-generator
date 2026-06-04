import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn
from models.schemas import StoryboardOutput, SlideStoryboard

TEMPLATE_DIR = os.getenv("TEMPLATE_DIR", "templates")
OUTPUT_DIR = os.getenv("OUTPUT_DIR", "outputs")


print(TEMPLATE_DIR, OUTPUT_DIR)

TEMPLATE_MAP = {
    "corporate":      "corporate.pptx",
    "sales":          "sales.pptx",
    "project_update": "project_update.pptx",
}

SLIDE_LAYOUT_MAP = {
    "title":   0,
    "content": 1,
    "data":    1,
    "summary": 1,
    "closing": 2,
}

SHAPE_TITLE    = "Title 1"
SHAPE_BODY     = "Content Placeholder 2"
SHAPE_SUBTITLE = "Subtitle 2"


class RenderingAgent:
    def __init__(self):
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    def _get_shape(self, slide, name: str):
        for shape in slide.shapes:
            print(shape.name)
            if shape.name == name and shape.has_text_frame:
                return shape
        return None

    def _set_text(self, shape, text: str):
        if shape is None:
            return
        shape.text_frame.clear()
        shape.text_frame.paragraphs[0].text = text

    def _set_bullets(self, shape, bullets: list[str]):
        if shape is None:
            return
        tf = shape.text_frame
        tf.clear()
        for i, bullet in enumerate(bullets[:5]):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = bullet
            p.level = 0

    def _add_image_placeholder(self, slide, description: str):
        left   = Inches(1.0)
        top    = Inches(4.5)
        width  = Inches(8.0)
        height = Inches(1.5)

        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"Image showing: {description}"
        p.alignment = PP_ALIGN.CENTER

        run = p.runs[0]
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        # Dashed border
        sp_pr = box._element.spPr
        ln = etree.SubElement(sp_pr, qn("a:ln"))
        ln.set("w", "12700")
        dash = etree.SubElement(ln, qn("a:prstDash"))
        dash.set("val", "dash")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(solid, qn("a:srgbClr"))
        clr.set("val", "AAAAAA")

    def _write_speaker_notes(self, slide, text: str):
        slide.notes_slide.notes_text_frame.text = text

    def _populate_title_slide(self, slide, data: SlideStoryboard, prs_title: str):
        self._set_text(self._get_shape(slide, SHAPE_TITLE), prs_title)
        self._set_text(self._get_shape(slide, SHAPE_SUBTITLE), data.key_message)
        self._write_speaker_notes(slide, data.speaker_notes)

    def _populate_content_slide(self, slide, data: SlideStoryboard):
        self._set_text(self._get_shape(slide, SHAPE_TITLE), data.slide_title)
        self._set_bullets(self._get_shape(slide, SHAPE_BODY), data.bullet_points)
        if data.image_required and data.image_description:
            self._add_image_placeholder(slide, data.image_description)
        self._write_speaker_notes(slide, data.speaker_notes)
    
    def _add_slide_number(self, slide, index: int, total: int):
        left   = Inches(8.5)
        top    = Inches(6.9)
        width  = Inches(1.2)
        height = Inches(0.3)

        box = slide.shapes.add_textbox(left, top, width, height)
        tf  = box.text_frame
        p   = tf.paragraphs[0]
        p.text = f"{index} / {total}"
        p.alignment = PP_ALIGN.RIGHT

        run = p.runs[0]
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    def run(self, job_id: str, template_name: str, storyboard: StoryboardOutput) -> str:
        template_file = TEMPLATE_MAP.get(template_name, "corporate.pptx")
        template_path = os.path.join(TEMPLATE_DIR, template_file)

        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template not found: {template_path}")

        prs = Presentation(template_path)

        # Safe slide removal — only remove slide relationships, not master
        XML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
        slide_ids = list(prs.slides._sldIdLst)
        for sldId in slide_ids:
            rId = sldId.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if rId:
                prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(sldId)

        for slide_data in storyboard.slides:
            layout_idx = SLIDE_LAYOUT_MAP.get(slide_data.slide_type, 1)
            layout = prs.slide_layouts[layout_idx]
            slide = prs.slides.add_slide(layout)

            if slide_data.slide_type == "title":
                self._populate_title_slide(slide, slide_data, storyboard.presentation_title)
            else:
                self._populate_content_slide(slide, slide_data)

        output_path = os.path.join(OUTPUT_DIR, f"{job_id}.pptx")
        prs.save(output_path)
        return output_path